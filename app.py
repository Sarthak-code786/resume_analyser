import json
import os
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, render_template_string, request
from groq import Groq
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document

load_dotenv()

app = Flask(__name__)

api_key = os.getenv("GROQ_API_KEY")
if not api_key:
    raise ValueError("GROQ_API_KEY environment variable is not set.")

client = Groq(api_key=api_key)
MODEL = "openai/gpt-oss-120b"
ROLE = "user"

job_description = """
Description
READY TO DO SOMETHING BIG?

How would you build the next generation of mission critical systems and services that power the ability to send, store and manage billions of products across the globe on behalf of our selling partners?

Fulfillment by Amazon (FBA) is the Amazon service for world-wide sellers that offers access to Amazon’s vast distribution network as a service. Similar to how AWS provides infrastructure as a service, FBA allows sellers to only pay for what they use, creating a seamless experience to leverage our world-class infrastructure. Sellers, from Fortune 50 companies to hundreds of thousands of small businesses and entrepreneurs, in 135 countries worldwide, rely upon our world-leading technology to grow and thrive on Amazon’s Marketplaces.

As a leader within FBA, you will provide both strategic and tactical oversight to build the features and functionality required to have the best in class offering for sellers. This leader will create mechanisms to drive optimal seller behavior for FBA inventory management, reducing total supply chain costs and making it easier for sellers to get their inventory close to customers, delivering more items with a Prime delivery promise. This leader will invent new ways to deliver value for sellers and their products, while leveraging the scale and innovation of technology from across the Amazon family, including AWS.

We operate as a startup within Amazon, providing our team great flexibility in product and technology direction. Come, join us and build the future!

Key job responsibilities
• You will solve ambiguous business problems and map them to elegant technical solutions
• You will use object-oriented concepts, algorithms and data structures to build systems at Amazon scale
• You will design, launch, and re-architect multi-tier systems with large amounts of data
• You will quickly adapt to fast-paced development environment with bias towards simple and extensible code

Basic Qualifications
- 3+ years of non-internship professional software development experience
- 2+ years of non-internship design or architecture (design patterns, reliability and scaling) of new and existing systems experience
- 3+ years of Video Games Industry (supporting title Development, Release, or Live Ops) experience
- Experience programming with at least one software programming language

Preferred Qualifications
- 3+ years of full software development life cycle, including coding standards, code reviews, source control management, build processes, testing, and operations experience
- Bachelor's degree in computer science or equivalent
"""


class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


class MatchResult(BaseModel):
    score: float
    details: dict


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    education: list[str] = []
    experiences: list[Experience] = []
    skills: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []
    total_experience_years: float | None = None


job_schema = JobD.model_json_schema()

system_prompt = f"""
You are an expert HR assistant.

Your job is to analyse the job descriptions and extract
structured information from it. You will return only valid JSON that matches the schema exactly:

{job_schema}
IMPORTANT:
Do NOT return the schema as it is.
Do NOT return fields like "properties", "type" or "title" in the output.
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

user_prompt = f"""
Analyse the following job description
and extract structured information from it. Return only the JSON object that matches the schema exactly.
{job_description}
"""

job_messages = [
    {"role": "system", "content": system_prompt},
    {"role": ROLE, "content": user_prompt},
]
job_response = client.chat.completions.create(
    model=MODEL,
    messages=job_messages,
    response_format={"type": "json_object"},
)
job_data = json.loads(job_response.choices[0].message.content)
job = JobD(**job_data)

resume_schema = Resume.model_json_schema()


def parse_resume(resume_text):
    system_prompt_for_resume = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning
    not only based on extracted section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships experience

    These may all contain relevant experience information. You should extract the information based on its meaning, not just the heading.

    Skills may also appear in the skills section, work experience,
    internships or projects.
    Return ONLY Valid JSON that matches the schema exactly:
    {resume_schema}

    IMPORTANT:
    1. Do NOT return the schema as it is.
    2. Do Not invent information.
    3. If a value is not available, return null
    4. If a list has no information, return an empty list.
    5. Include internships inside experiences.
    6. Extract skills from all sections, not just the skills section.
    """

    user_prompt_for_resume = f"""
    Parse the following resume:
    {resume_text}
    """

    messages = [
        {"role": "system", "content": system_prompt_for_resume},
        {"role": ROLE, "content": user_prompt_for_resume},
    ]
    response = client.chat.completions.create(
        model=MODEL,
        messages=messages,
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return Resume(**data)


def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:
    {match_schema}

    Give me:
    1. Candidate name
    2. Matching skills
    3. Missing skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """

    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)


def read_pdf(file_path):
    reader = PdfReader(str(file_path))
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx(file_path):
    doc = Document(str(file_path))
    text = ""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(file_path)
    if suffix == ".docx":
        return read_docx(file_path)
    raise ValueError("Unsupported file type")


HTML_PAGE = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>Resume Matcher</title>
  <style>
    body { font-family: Arial, sans-serif; margin: 40px; background: #f4f7fb; color: #1f2937; }
    .card { max-width: 850px; margin: 0 auto; background: white; padding: 30px; border-radius: 16px; box-shadow: 0 12px 35px rgba(0,0,0,0.08); }
    input[type=file] { margin-top: 12px; }
    button { background: #2563eb; color: white; border: none; padding: 12px 20px; border-radius: 8px; cursor: pointer; font-size: 16px; }
    .result { margin-top: 24px; padding: 18px; border-radius: 10px; background: #f8fafc; }
    h2 { margin-top: 8px; }
    code { background: #eef2ff; padding: 2px 6px; border-radius: 6px; }
  </style>
</head>
<body>
  <div class="card">
    <h1>Resume Match Checker</h1>
    <p>Upload a resume and compare it against the target Amazon-style role.</p>
    <form method="post" enctype="multipart/form-data">
      <input type="file" name="resume" accept=".pdf,.docx">
      <br><br>
      <button type="submit">Analyze Resume</button>
    </form>

    {% if result %}
      <div class="result">
        <h2>Result</h2>
        <p><strong>Name:</strong> {{ result.name }}</p>
        <p><strong>Overall Match:</strong> {{ result.score }}%</p>
        <p><strong>Matching Skills:</strong> {{ result.matching_skills }}</p>
        <p><strong>Missing Skills:</strong> {{ result.missing_skills }}</p>
        <p><strong>Experience Requirement Met:</strong> {{ result.experience_requirement_met }}</p>
        <p><strong>Verdict:</strong> {{ result.verdict }}</p>
      </div>
    {% endif %}

    {% if error %}
      <div class="result" style="background: #fff1f2; color: #b91c1c;">
        <strong>Error:</strong> {{ error }}
      </div>
    {% endif %}
  </div>
</body>
</html>
"""


@app.get("/")
def home():
    return render_template_string(HTML_PAGE, result=None, error=None)


@app.post("/")
def upload_resume():
    file = request.files.get("resume")
    if not file or not file.filename:
        return render_template_string(HTML_PAGE, result=None, error="Please upload a PDF or DOCX resume."), 400

    try:
        upload_path = Path("uploads")
        upload_path.mkdir(exist_ok=True)
        saved_path = upload_path / file.filename
        file.save(saved_path)

        raw_text = read_resume(saved_path)
        parsed_resume = parse_resume(raw_text)
        result = final_score(job, parsed_resume)

        score = float(result.score)
        details = result.details
        response = {
            "name": details.get("name") or details.get("candidate_name") or parsed_resume.name,
            "score": score,
            "matching_skills": details.get("matching_skills", []),
            "missing_skills": details.get("missing_skills", []),
            "experience_requirement_met": details.get("experience_requirement_met", False),
            "verdict": details.get("verdict") or details.get("final_verdict") or "No verdict provided.",
        }

        return render_template_string(HTML_PAGE, result=response, error=None)
    except Exception as exc:
        return render_template_string(HTML_PAGE, result=None, error=f"Processing failed: {exc}"), 500


if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=False)
