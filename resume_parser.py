import json
import os
import time
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel

# structured prompt for the model

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("GROQ_API_KEY environment variable is not set.")

client = Groq(api_key=my_api_key)
model = "openai/gpt-oss-120b"
role = "user"

job_description = """
Description
READY TO DO SOMETHING BIG?

How would you build the next generation of mission critical systems and services that power the ability to send, store and manage billions of products across the globe on behalf of our selling partners?

Fulfillment by Amazon (FBA) is the Amazon service for world-wide sellers that offers access to Amazon’s vast distribution network as a service. Similar to how AWS provides infrastructure as a service, FBA allows sellers to only pay for what they use, creating a seamless experience to leverage our world-class infrastructure. Sellers, from Fortune 50 companies to hundreds of thousands of small businesses and entrepreneurs, in 135 countries worldwide, rely upon our world-leading technology to grow and thrive on Amazon’s Marketplaces.

As a leader within FBA, you will provide both strategic and tactical oversight to build the features and functionality required to have the best in class offering for sellers. This leader will create mechanisms to drive optimal seller behavior for FBA inventory management, reducing total supply chain costs and making it easier for sellers to get their inventory close to customers, delivering more items with a Prime delivery promise. This leader will invent new ways to deliver value for sellers and their products, while leveraging the scale and innovation of technology from across the Amazon family, including AWS.

We operate as a startup within Amazon, providing our team great flexibility in product and technology direction. Come, join us and build the future!

Key job responsibilities
• You will solve ambiguous business problems and map them to elegant technical solutions
• You will use object-oriented concepts, algorithms and data structures to build systems at Amazon scale
• You will design, launch, and re-architect multi-tier systems with large amounts of data
• You will quickly adapt to fast-paced development environment with bias towards simple and extensible code

Basic Qualifications
- 3+ years of non-internship professional software development experience
- 2+ years of non-internship design or architecture (design patterns, reliability and scaling) of new and existing systems experience
- 3+ years of Video Games Industry (supporting title Development, Release, or Live Ops) experience
- Experience programming with at least one software programming language

Preferred Qualifications
- 3+ years of full software development life cycle, including coding standards, code reviews, source control management, build processes, testing, and operations experience
- Bachelor's degree in computer science or equivalent
"""

class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]

job_schema = JobD.model_json_schema()

system_prompt = f"""
    "You are an expert HR assistant.

    Your job is to analyse the job descriptions and extract
    structured information from it. You will return only valid JSON that matches the schema exactly:

    {job_schema}
IMPORTANT:
Do NOT return the schema as it is.
Do NOT return fields like "properties", "type" or "title" in the output.
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

user_prompt = f"""
Analyse the following job description
and extract structured information from it. Return only the JSON object that matches the schema exactly.
{job_description}
"""
message_system={
    "role": "system",
    "content": system_prompt,
    }
message_user={
    "role": role,
    "content": user_prompt,
    }
response_format={
    "type": "json_object",
}

messages = [message_system, message_user]
response = client.chat.completions.create(
    model=model,
    messages=messages,
    response_format=response_format
)

answer = response.choices[0].message.content
raw_json = answer
job_data = json.loads(raw_json)
job = JobD(**job_data)

print(job.minimum_experience)
print(job.education_requirements)

class MatchResult(BaseModel):
    score: float
    details: dict
class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills: list[str]=[]

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    education: list[str]=[]
    experiences: list[Experience]=[]
    skills: list[str]=[]
    projects: list[str]=[]
    certifications: list[str]=[]

    total_experience_years: float | None = None
resume_schema = Resume.model_json_schema()
def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    you are an HR recruiter.

    Compare the candidate's resume with the job description

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:
    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing skills
    4. Whether experience requirement is met
    5. Overall match percentage froma 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """

    message={
        "role": "user",
        "content": prompt
    }
    message=[message]
    response_format={
        "type": "json_object",
    }
    response = client.chat.completions.create(model=model, messages=message, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    System_prompt = f"""
    You are an expert resume parser.

    Extract information from the reusme nased on its meaning
    not only based on extracted section headings

    Different resumes may use differnet headings

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships experience

    These may all contain relevant experience information. You should extract the information based on its meaning, not just the heading.

    Skills may also appear in the skills section, work experience,
    internships or projects.
    Return ONLY Valid JSON that matches the schema exactly:
    {resume_schema}

    IMPORTANT:
    1. Do NOT return the schema as it is.
    2. Do Not invent information.
    3. If a value is not available, return null
    4. If a list has no information, return an empty list.
    5. Include internships inside experiences.
    6. Extract skills from all sections, not just the skills section.
    """

    user_prompt = f"""
    Parse the following resume:
    {resume_text}
    """
    message_system={
        "role": "system",
        "content": System_prompt,
    }
    message_user={
        "role": role,
        "content": user_prompt,
    }
    messages = [message_system, message_user]
    response_format={
        "type": "json_object",
    }
    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )
    answer = response.choices[0].message.content
    raw_output = answer
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume

from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    doc = Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


resume_folder = Path(__file__).resolve().parent / "Resumes"
all_results = []
for file_path in sorted(resume_folder.iterdir()):
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume = parse_resume(resume_text)
    time.sleep(5)
    result = final_score(job, parsed_resume)
    time.sleep(5)
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
all_results.sort(
    key=lambda candidate: candidate["score"], 
    reverse=True #descending order sort
)
top_2 = all_results[:2]
worst_2 = all_results[-2:]

print("\nTop 2 Candidates:")
for candidate in top_2:
    print(f"Name: {candidate['name']}, Score: {candidate['score']}")
    print("Details:", candidate["details"])
    print()

print("\nWorst 2 Candidates:")
for candidate in worst_2:
    print(f"Name: {candidate['name']}, Score: {candidate['score']}")
    print("Details:", candidate["details"])
    print()
    
